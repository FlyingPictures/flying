from pathlib import Path
import json, re, shutil, html
from urllib.parse import quote
import openpyxl

ROOT = Path(__file__).resolve().parent
SITE = ROOT / 'site'
ASSETS = SITE / 'assets'
SOURCE = Path('/Users/enriquevelasco/Downloads/Paquetes_FlyingPictures_Teotihuacan_Sean_09092026.xlsx')
sheet = openpyxl.load_workbook(SOURCE, data_only=True)['Paquetes']
esc = html.escape
T = lambda es, en: {'es': es, 'en': en}
names = ['essential', 'pyramids', 'premium', 'celebration', 'traveler', 'private', 'charter']
english = ['Essential Flight', 'Flight + Transport + Pyramids', 'Premium – VIP', 'Celebration', 'Traveler', 'Private Balloon Experience', 'FPM Private Charter']
descriptions = [
 T('El amanecer, el valle y la emoción de volar. La experiencia esencial para quienes quieren disfrutar del globo a su manera.', 'The sunrise, the valley and the joy of flying. The essential experience for those who want to enjoy a balloon flight their own way.'),
 T('Del vuelo al encuentro con las pirámides. Combina tu aventura en el aire con desayuno, transporte y acceso a la zona arqueológica.', 'From the sky to the pyramids. Pair your balloon flight with breakfast, transportation and admission to the archaeological site.'),
 T('Una experiencia privada por aire y tierra: champagne, brunch en Mayahuel, pirámides y un recorrido por cuevas ceremoniales.', 'A private experience in the air and on the ground: champagne, brunch at Mayahuel, pyramids and a ceremonial cave tour.'),
 T('Un vuelo privado para una ocasión que merece algo extraordinario. Flores, desayuno y detalles para celebrar juntos.', 'A private flight for an occasion that deserves something extraordinary. Flowers, breakfast and thoughtful touches to celebrate together.'),
 T('Para quienes quieren descubrir más. Vuelo, pirámides con guía bilingüe, artesanías y bebidas típicas en una misma experiencia.', 'For those who want to discover more. A balloon flight, pyramids with a bilingual guide, local crafts and traditional drinks in one experience.'),
 T('Una canastilla solo para tu grupo. Comparte el cielo con quienes tú elijas y celebra con champagne y desayuno.', 'A basket exclusively for your group. Share the sky with the people you choose, followed by champagne and breakfast.'),
 T('Tu experiencia más completa: vuelo privado, transporte premium, guía bilingüe desde la recogida y brunch en Club Campestre.', 'Your most complete experience: a private flight, premium transportation, a bilingual guide from pickup and brunch at Club Campestre.')]
features_en = {21:'Pre-flight coffee break',22:'Flight insurance',23:'Traditional sparkling wine toast',24:'Champagne toast',25:'Flight certificate',26:'Light buffet breakfast before takeoff',27:'Full breakfast',28:'Extra breakfast',29:'Shared round-trip transportation',30:'Private transportation',31:'Private round-trip transportation',32:'Premium private round-trip transportation',33:'Upgrade to a private flight',34:'Personalized banner',35:'Flower bouquet',36:'Mariachi music upon landing',37:'Champagne',38:'Cake and/or birthday song',39:'Admission to the archaeological site',40:'Bilingual guide',41:'Bilingual guide from pickup',42:'Ceremonial cave tour',43:'Photos and video',44:'Traveler souvenir kit',45:'Waiter service',46:'Canapés and snacks',47:'Brunch at Mayahuel',48:'Brunch at Club Campestre',49:'Local crafts and traditional drinks workshop'}
features_es = {21:'Café y refrigerio antes del vuelo',22:'Seguro de vuelo',23:'Brindis tradicional con vino espumoso',24:'Brindis con champagne',25:'Certificado de vuelo',26:'Desayuno buffet ligero antes del despegue',27:'Desayuno completo',28:'Desayuno extra',29:'Transporte redondo compartido',30:'Transporte privado',31:'Transporte redondo privado',32:'Transporte redondo privado premium',33:'Conversión a vuelo privado',34:'Lona personalizada',35:'Ramo de flores',36:'Mariachis al aterrizar',37:'Champagne',38:'Pastel y/o mañanitas',39:'Acceso a la zona arqueológica',40:'Guía bilingüe',41:'Guía bilingüe desde la recogida',42:'Tour en cuevas ceremoniales',43:'Fotografías y video',44:'Kit de recuerdos Viajero',45:'Servicio de mesero',46:'Canapés y snacks',47:'Brunch en Mayahuel',48:'Brunch en Club Campestre',49:'Taller de artesanías y bebidas típicas'}
extra_translations = [
 ('Transporte Redondo Privado 1-4Pax , Zona Centrica', 'Private round-trip transportation, 1–4 guests, central area'),
 ('Transporte Redondo Privado 2 Pax , Zona Centrica', 'Private round-trip transportation, 2 guests, central area'),
 ('Guía Bilingüe (1-6 Personas)', 'Bilingual guide (1–6 guests)'),
 ('Servicio por acompañante extra (no vuela)', 'Additional non-flying companion'),
 ('Persona Adicional Que no Vuela (VIP)', 'Additional non-flying guest (VIP)'),
 ('Persona Adicional Adulto VIP', 'Additional adult (VIP)'),('Persona Adicional Niño VIP','Additional child (VIP)'),
 ('Persona Adicional (Celebración)', 'Additional guest (Celebration)'),('Persona Adicional (Viajero)', 'Additional guest (Traveler)'),
 ('Persona Adicional (Paq 6)', 'Additional guest (Private Balloon Experience)'),('Persona Adicional (FPM)', 'Additional guest (FPM)'),
 ('Desayuno Completo para acompañante (Buffette)', 'Full buffet breakfast for a companion'),
 ('Fotografías y Video', 'Photos and video'),('Mariachis al aterrizaje', 'Mariachi music upon landing'),('Ramo de Flores', 'Flower bouquet'),
 ('Acceso a la Zona Arqueológica', 'Admission to the archaeological site'),('Pastel (5 Rebanadas)', 'Cake (5 slices)'),
 ('Lona Estandar', 'Standard banner'),('Feliz Cumpleaños', 'Happy Birthday'),('Te Amo', 'I Love You'),('Feliz Aniversario', 'Happy Anniversary'),
 ('¿Te Quieres Casar Conmigo?', 'Will You Marry Me?'),('¿Quieres ser mi novia?', 'Will You Be My Girlfriend?')]
def translate_extra(value):
    for a,b in extra_translations: value=value.replace(a,b)
    return value

packages=[]
for i,slug in enumerate(names):
    col=i+2
    includes=[]
    for row in range(21,50):
        if sheet.cell(row,col).value is True:
            # The champagne row repeats the champagne toast in these packages.
            if row==37 and sheet.cell(24,col).value: continue
            includes.append(T(features_es[row],features_en[row]))
    extras=[T(str(sheet.cell(r,col).value).strip(),translate_extra(str(sheet.cell(r,col).value).strip())) for r in range(51,64) if sheet.cell(r,col).value]
    packages.append(dict(id=slug,name=T(sheet.cell(2,col).value,english[i]),description=descriptions[i],type='shared' if sheet.cell(3,col).value=='Compartido' else 'private',price=sheet.cell(11,col).value or sheet.cell(13,col).value,child=sheet.cell(12,col).value,unit='person' if sheet.cell(11,col).value else 'package',includes=includes,extras=extras,bookingUrl=None))

# Optional verified booking URLs. Until provided, every enquiry uses WhatsApp.
config=json.loads((ROOT/'booking-links.json').read_text())
for p in packages: p['bookingUrl']=config.get(p['id'])
(ROOT/'review'/'package-data.json').write_text(json.dumps(packages,ensure_ascii=False,indent=2))

copy={
 'nav':T('Experiencias','Experiences'),'help':T('Te ayudamos a elegir','Let us help you choose'),'explore':T('Explorar vuelos','Explore flights'),
 'eyebrow':T('TEOTIHUACÁN, MÉXICO','TEOTIHUACÁN, MEXICO'),
 'title':T('Hay amaneceres<br>que cambian todo.','Some sunrises<br>change everything.'),
 'subtitle':T('El tuyo empieza en el cielo.','Yours begins in the sky.'),
 'intro':T('Vuela sobre el valle de Teotihuacán. Encuentra la experiencia que va contigo y haz de esta mañana un recuerdo extraordinario.','Fly above the Teotihuacán Valley. Find the experience that feels like you, and make this morning an extraordinary memory.'),
 'heroCta':T('Encuentra tu vuelo','Find your flight'),'heroOther':T('Habla con nosotros','Talk to us'),
 'shared':T('Vuelos compartidos','Shared flights'),'private':T('Vuelos privados','Private flights'),'all':T('Todos los vuelos','All flights'),
 'sectionEye':T('SIETE FORMAS DE VIVIRLO','SEVEN WAYS TO EXPERIENCE IT'),
 'sectionTitle':T('Tu mañana. Tu forma de volar.','Your morning. Your way to fly.'),
 'sectionBody':T('Desde la emoción de tu primer vuelo hasta una celebración privada. Compara, descubre los detalles y elige tu experiencia.','From the thrill of your first flight to a private celebration. Compare, explore the details and find your experience.'),
 'details':T('Ver experiencia','Explore experience'),'included':T('Tu experiencia incluye','Your experience includes'),
 'extras':T('Hazla aún más tuya','Make it your own'),'extrasBody':T('Servicios opcionales con costo adicional. Consulta disponibilidad y precio al reservar.','Optional services at an additional cost. Ask about availability and pricing when booking.'),
 'price':T('Precio de la experiencia','Experience price'),'person':T('por adulto','per adult'),'package':T('por paquete','per package'),
 'child':T('Niños de 4 a 12 años','Children aged 4–12'),'from':T('Desde','From'),
 'ask':T('Consultar disponibilidad','Check availability'),'book':T('Reservar experiencia','Book experience'),
 'via':T('Te atendemos por WhatsApp','Connect with us on WhatsApp'),
 'groupNote':T('Consulta con nuestro equipo el número de personas incluidas y el total para tu grupo.','Ask our team how many guests are included and confirm the total for your group.'),
 'twoNote':T('Vuelo privado con mínimo de dos personas. Confirma las personas incluidas y el total para tu grupo.','Private flight with a minimum of two guests. Confirm the included guests and the total for your group.'),
 'priceNote':T('Precios en pesos mexicanos (MXN). Confirma disponibilidad, condiciones y total antes de reservar.','Prices in Mexican pesos (MXN). Confirm availability, terms and the total before booking.'),
 'fullPage':T('Abrir página del paquete','Open package page'),'close':T('Cerrar detalles','Close details'),'back':T('Ver todas las experiencias','View all experiences'),
 'faqTitle':T('Antes de despegar','Before you take off'),'faqEye':T('TODO EMPIEZA CON UNA BUENA ELECCIÓN','A GREAT EXPERIENCE STARTS HERE'),
 'contactTitle':T('Una mañana extraordinaria<br>empieza con un hola.','An extraordinary morning<br>starts with hello.'),
 'contactBody':T('Cuéntanos cuándo quieres volar, cuántas personas vienen y qué te gustaría vivir. Te ayudamos a encontrar tu experiencia.','Tell us when you would like to fly, how many guests are joining and what you have in mind. We will help you find your experience.'),
 'contactCta':T('Planear mi vuelo por WhatsApp','Plan my flight on WhatsApp'),
 'footer':T('AN AVIATION COMPANY','AN AVIATION COMPANY'),
 'skip':T('Ir al contenido','Skip to content'),'filterLabel':T('Filtrar experiencias','Filter experiences'),
 'count':T('experiencias para descubrir','experiences to discover'),'tip':T('Una experiencia, a tu medida.','An experience, made yours.'),
 'call':T('Llámanos','Call us'),'currency':T('Todos los precios en MXN','All prices in MXN')}
faq=[
 (T('¿Cuál es la diferencia entre compartido y privado?','What is the difference between shared and private?'),T('En un vuelo compartido viajas en la canastilla con otros pasajeros. En un vuelo privado, la canastilla es exclusiva para tu grupo y el piloto. Puedes filtrar las experiencias para comparar ambas opciones.','On a shared flight, you share the basket with other passengers. On a private flight, the basket is exclusively for your group and the pilot. Filter the experiences to compare both options.')),
 (T('¿Todos los paquetes incluyen transporte y desayuno?','Do all packages include transportation and breakfast?'),T('No. Vuelo Esencial no incluye desayuno completo ni transporte. Los demás paquetes tienen distintas inclusiones: abre los detalles para revisar exactamente qué incluye tu elección.','No. Essential Flight does not include a full breakfast or transportation. Other packages have different inclusions: open the details to see exactly what is included in your choice.')),
 (T('¿El precio es por persona o por paquete?','Is the price per person or per package?'),T('Vuelo Esencial y +Transporte + Pirámides muestran precio por adulto y tarifa infantil. Las otras experiencias muestran precio por paquete. Nuestro equipo te confirma las personas incluidas y el total de tu reserva.','Essential Flight and Flight + Transport + Pyramids show adult and child rates. The other experiences show a package price. Our team will confirm the included guests and your booking total.')),
 (T('¿Cómo consulto fechas o reservo?','How do I check dates or book?'),T('El botón de disponibilidad abre una conversación con el nombre de tu paquete. Comparte tu fecha, número de adultos y niños y cualquier solicitud especial. Nuestro equipo confirmará disponibilidad y los siguientes pasos.','The availability button opens a conversation with your package name. Share your date, number of adults and children and any special requests. Our team will confirm availability and the next steps.')),
 (T('¿Qué pasa si el clima no permite volar?','What happens if the weather does not allow a flight?'),T('Los vuelos dependen de las condiciones meteorológicas. Antes de reservar, consulta con nuestro equipo las condiciones de cambio de fecha, cancelación y reembolso aplicables a tu experiencia.','Flights depend on weather conditions. Before booking, ask our team about the rescheduling, cancellation and refund terms that apply to your experience.'))]

def whatsapp(lang,p=None):
    message=(f'Hola, me interesa el paquete {p["name"][lang]}. ¿Me ayudan con disponibilidad, personas incluidas y precio total? Mi fecha es: ' if lang=='es' else f'Hello, I am interested in the {p["name"][lang]} package. Could you help me with availability, included guests and the total price? My preferred date is: ') if p else ('Hola, me gustaría planear un vuelo en globo en Teotihuacán. ¿Me ayudan a elegir un paquete?' if lang=='es' else 'Hello, I would like to plan a balloon flight in Teotihuacán. Could you help me choose a package?')
    return 'https://wa.me/525510484758?text='+quote(message)
def money(n): return '${:,.0f}'.format(n)
def a(text,url,cls='button',attrs=''): return f'<a class="{cls}" href="{esc(url)}" {attrs}>{text}</a>'
def ext(text,url,cls='button'): return a(text,url,cls,'target="_blank" rel="noopener noreferrer"')
def bullets(items,lang,cls=''): return f'<ul class="{cls}">'+''.join(f'<li>{esc(v[lang])}</li>' for v in items)+'</ul>'
def body_detail(p,lang,base,heading='h2'):
    c=lambda k:copy[k][lang]
    note=c('twoNote') if p['id'] in ['celebration','private'] else c('groupNote')
    child=f'<p class="child-price">{c("child")}: <strong>{money(p["child"])} MXN</strong></p>' if p['child'] else ''
    extra=bullets(p['extras'],lang,'extras-list')
    exclusions=[]
    if not any(v['en'] in ['Shared round-trip transportation','Private round-trip transportation','Premium private round-trip transportation'] for v in p['includes']): exclusions.append(T('Transporte no incluido en el precio base.','Transportation is not included in the base price.'))
    if p['id']=='essential': exclusions.append(T('Desayuno completo y acceso a las pirámides no incluidos.','Full breakfast and pyramid admission are not included.'))
    if p['id']=='pyramids': exclusions.append(T('Guía bilingüe disponible con costo adicional.','Bilingual guide available at an additional cost.'))
    if p['id'] in ['pyramids','traveler']:
        private_charge=4999 if p['id']=='pyramids' else 5999
        extras_note=T(f'Conversión a vuelo privado: +{money(private_charge)} MXN. Transporte privado: +$2,000 MXN. Consulta condiciones y total para tu grupo.',f'Private flight upgrade: +{money(private_charge)} MXN. Private transportation: +$2,000 MXN. Ask about terms and the total for your group.')
        extra+=f'<p class="note">{extras_note[lang]}</p>'
    return f'''<div class="detail-photo"><img src="{base}assets/{p['id']}.webp" alt="{esc(p['name'][lang])}" width="850" height="560"><span class="tag">{c(p['type'])}</span></div>
    <div class="detail-content"><p class="eyebrow">{c('eyebrow')}</p><{heading} class="detail-title">{esc(p['name'][lang])}</{heading}><p class="lead">{p['description'][lang]}</p>
    <div class="detail-price"><span>{c('price')}</span><div><strong>{money(p['price'])}</strong> <span>MXN · {c(p['unit'])}</span></div>{child}<p class="small">{note if p['unit']=='package' else c('priceNote')}</p></div>
    <h3>{c('included')}</h3>{bullets([T('Vuelo en globo '+('compartido' if p['type']=='shared' else 'privado'),('Shared' if p['type']=='shared' else 'Private')+' hot air balloon flight')]+p['includes'],lang,'included-list')}
    {''.join(f'<p class="note">{e[lang]}</p>' for e in exclusions)}
    <details class="extras"><summary>{c('extras')}</summary><p class="small">{c('extrasBody')}</p>{extra}</details>
    <p class="small terms">{c('priceNote')}</p></div>
    <div class="detail-bottom"><div><strong>{money(p['price'])} <small>MXN</small></strong><span>{c(p['unit'])}</span></div>{ext(c('book' if p['bookingUrl'] else 'ask')+' ↗',p['bookingUrl'] or whatsapp(lang,p),'button amber')}</div>'''

def render(lang,p=None):
    c=lambda k:copy[k][lang]
    # index.html / en/index.html / packages/id.html / en/packages/id.html
    depth=(1 if lang=='en' else 0)+(1 if p else 0)
    base='../'*depth
    home=('../' if p else '')+'index.html'
    other='en' if lang=='es' else 'es'
    other_path=base+('en/' if other=='en' else '')+(f'packages/{p["id"]}.html' if p else 'index.html')
    title=(p['name'][lang] if p else ('Vuelos en globo en Teotihuacán' if lang=='es' else 'Teotihuacán hot air balloon flights'))+' | Flying Pictures México'
    desc=p['description'][lang] if p else c('intro')
    current_path=home if not p else p['id']+'.html'
    languages=''.join(f'<a data-language="{l}" href="{current_path if l==lang else other_path}?lang={l}" lang="{l}" '+('aria-current="true" ' if l==lang else '')+f'aria-label="{"Ver en español" if l=="es" else "View in English"}">{l.upper()}</a>' for l in ['es','en'])
    nav=f'''<a class="skip" href="#main">{c('skip')}</a><header class="header"><nav class="nav" aria-label="{'Navegación principal' if lang=='es' else 'Main navigation'}"><a class="brand" href="{home}" aria-label="Flying Pictures México — {'Inicio' if lang=='es' else 'Home'}"><img src="{base}assets/logo.png" alt="Flying Pictures México" width="60" height="80"></a><div class="nav-links"><a href="{home}#experiences">{c('nav')}</a><a href="{home}#faq">{c('faqTitle')}</a></div><div class="nav-right"><div class="language" aria-label="{'Idioma' if lang=='es' else 'Language'}">{languages}</div>{a(c('explore'),home+'#experiences','button amber nav-cta')}</div></nav></header>'''
    footer=f'''<footer><div class="footer-brand">Flying Pictures<span>MÉXICO</span></div><p>{c('footer')}</p><div><a href="tel:+525580251057">{c('call')}: +52 55 8025 1057</a><a href="mailto:volamoscontigo@flyingpictures.mx">volamoscontigo@flyingpictures.mx</a></div><p class="copyright">© 2026 Flying Pictures México · Teotihuacán</p></footer>'''
    if p:
        content=f'<main id="main" class="package-page"><a class="back-link" href="{home}#experiences">← {c("back")}</a><article class="standalone-detail">{body_detail(p,lang,base,"h1")}</article></main>'
    else:
        cards=[]; templates=[]
        for idx,q in enumerate(packages):
            url=f'packages/{q["id"]}.html'
            highlights=[x for x in q['includes'] if x['en'] not in ['Flight insurance','Flight certificate','Pre-flight coffee break']][:3]
            if len(highlights)<3: highlights=q['includes'][:3]
            cards.append(f'''<article class="package-card {'charter-card' if q['id']=='charter' else ''}" data-type="{q['type']}" id="{q['id']}"><a href="{url}" data-package="{q['id']}" class="card-image" aria-label="{c('details')}: {esc(q['name'][lang])}"><img src="{base}assets/{q['id']}.webp" alt="{esc(q['name'][lang])}" loading="lazy" width="850" height="560"><span class="tag">{c(q['type'])}</span><span class="image-arrow" aria-hidden="true">↗</span></a><div class="card-content"><span class="card-number">0{idx+1} / FLYING EXPERIENCES</span><h3><a href="{url}" data-package="{q['id']}">{esc(q['name'][lang])}</a></h3><p class="card-description">{q['description'][lang]}</p>{bullets(highlights,lang,'card-highlights')}<div class="card-foot"><div class="card-price"><strong>{money(q['price'])} <small>MXN</small></strong><span>{c(q['unit'])}</span></div>{a(c('details')+' <span aria-hidden="true">↗</span>',url,'details-link',f'data-package="{q["id"]}"')}</div></div></article>''')
            alternate=base+('en/' if other=='en' else '')+f'packages/{q["id"]}.html?lang={other}'
            templates.append(f'<template id="detail-{q["id"]}"><div class="dialog-language">{a("English" if other=="en" else "Español",alternate,"language-link",f"lang=\"{other}\"")}</div>{body_detail(q,lang,base)}<div class="permalink">{a(c("fullPage")+" ↗",url)}</div></template>')
        faq_html=''.join(f'<details class="faq-item"><summary>{x[lang]}</summary><p>{y[lang]}</p></details>' for x,y in faq)
        content=f'''<main id="main"><section class="hero"><img class="hero-image" src="{base}assets/hero.webp" alt="{'La emoción de volar en globo en Teotihuacán' if lang=='es' else 'The joy of ballooning in Teotihuacán'}" width="1800" height="1200" fetchpriority="high"><div class="hero-shade"></div><div class="hero-content"><p class="eyebrow">{c('eyebrow')} <span>✦</span> FLYING PICTURES</p><h1>{c('title')}</h1><p class="hero-serif">{c('subtitle')}</p><p class="hero-description">{c('intro')}</p><div class="hero-actions">{a(c('heroCta')+' ↓','#experiences','button amber')}{ext(c('heroOther')+' ↗',whatsapp(lang),'hero-secondary')}</div></div><div class="hero-caption"><span>19°41′ N &nbsp; 98°49′ W</span><span>{c('tip')}</span></div></section>
        <div class="promise-bar"><span>✦ &nbsp; {c('shared')}</span><span>✦ &nbsp; {c('private')}</span><span>✦ &nbsp; {'Atención en español e inglés' if lang=='es' else 'English & Spanish support'}</span></div>
        <section class="experiences section" id="experiences"><div class="section-heading"><p class="eyebrow">{c('sectionEye')}</p><h2>{c('sectionTitle')}</h2><p>{c('sectionBody')}</p></div><div class="filter-bar"><div class="filters" role="group" aria-label="{c('filterLabel')}" hidden>{''.join(f'<button type="button" data-filter="{k}" aria-pressed="{str(k=="all").lower()}">{c(k)}</button>' for k in ['all','shared','private'])}</div><span class="currency">{c('currency')}</span></div><p class="sr-only" id="filter-status" aria-live="polite" data-label="{c('count')}"></p><div class="packages-grid">{''.join(cards)}</div><p class="price-disclaimer">{c('priceNote')}</p></section>
        <section class="faq-section section" id="faq"><div><p class="eyebrow">{c('faqEye')}</p><h2>{c('faqTitle')}</h2><p>{c('help')}.</p>{ext('WhatsApp ↗',whatsapp(lang),'text-link')}</div><div>{faq_html}</div></section>
        <section class="contact-section" id="contact"><p class="eyebrow">FLYING PICTURES MÉXICO</p><h2>{c('contactTitle')}</h2><p>{c('contactBody')}</p>{ext(c('contactCta')+' ↗',whatsapp(lang),'button amber')}</section></main>
        <div class="mobile-bar">{a(c('nav'),'#experiences','button light')}{ext('WhatsApp ↗',whatsapp(lang),'button amber')}</div>
        <dialog id="package-dialog" aria-label="{c('details')}"><button class="dialog-close" type="button" aria-label="{c('close')}" autofocus>×</button><div id="dialog-body"></div></dialog>{''.join(templates)}'''
    return f'''<!doctype html><html lang="{lang}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="theme-color" content="#102c35"><title>{esc(title)}</title><meta name="description" content="{esc(desc)}"><meta property="og:title" content="{esc(title)}"><meta property="og:description" content="{esc(desc)}"><meta property="og:type" content="website"><link rel="icon" href="{base}assets/favicon.svg"><link rel="stylesheet" href="{base}assets/fonts.css"><link rel="stylesheet" href="{base}assets/style.css"><link rel="alternate" hreflang="{other}" href="{other_path}"><script defer src="{base}assets/app.js"></script></head><body>{nav}{content}{footer}</body></html>'''

for lang in ['es','en']:
    target=SITE/('en' if lang=='en' else '')
    (target/'packages').mkdir(parents=True,exist_ok=True)
    (target/'index.html').write_text(render(lang))
    for p in packages: (target/'packages'/f'{p["id"]}.html').write_text(render(lang,p))

# Reuse locally built brand fonts; no third-party font requests at runtime.
css=(ROOT.parent/'.next/static/chunks/e84ce45ac846b1b2.css').read_text()
font_css=[]
for face in re.findall(r'@font-face\{[^}]+\}',css):
    if 'unicode-range:U+??' not in face: continue
    if not any(f'font-family:{family};' in face for family in ['Inter','Poppins','Libre Baskerville']): continue
    name=re.search(r'../media/([^\)]+)',face).group(1)
    shutil.copy(ROOT.parent/'.next/static/media'/name,ASSETS/name)
    font_css.append(face.replace('../media/',''))
(ASSETS/'fonts.css').write_text('\n'.join(font_css))
shutil.copy(ROOT.parent/'public/favicon.svg',ASSETS/'favicon.svg')
print(f'Generated 16 static pages in {SITE}')
